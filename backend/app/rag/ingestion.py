"""文档/表格入库：文本抽取 -> 分块 -> 向量化 -> 落库；表格类数据写入 SQLite。"""
from __future__ import annotations

import uuid
from pathlib import Path

import numpy as np
import pandas as pd
from sqlalchemy.orm import Session

from app.config import Settings
from app.db.models import Chunk, Dataset, Document
from app.llm.embeddings import get_embedder
from app.rag.chunker import recursive_split
from app.rag.store import vector_store

MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".csv"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        # 优先 PyMuPDF：对内嵌 CID 字体/缺 ToUnicode 映射的中文 PDF（如港股年报）
        # 提取质量远好于 pypdf；pypdf 仅作兜底
        try:
            import pymupdf

            with pymupdf.open(str(path)) as doc:
                return "\n".join(page.get_text() for page in doc)
        except ImportError:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from docx import Document as Docx  # 可选依赖

        doc = Docx(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"暂不支持的文档格式: {suffix}")


EMBED_BATCH_SIZE = 16  # OpenAI 兼容 Embedding 接口的批量上限（SiliconFlow 为 32，留余量）


def ingest_document(
    db: Session, settings: Settings, dataset: Dataset, path: Path, filename: str
) -> Document:
    text = extract_text(path)
    if not text.strip():
        raise ValueError("文件内容为空，无法解析出文本")

    chunks = recursive_split(text, settings.rag_chunk_size, settings.rag_chunk_overlap)
    embedder = get_embedder(settings)
    # 分批向量化：大文档（数百页 PDF）一次性提交会超出 Embedding API 批量限制
    all_vecs: list[np.ndarray] = []
    for i in range(0, len(chunks), EMBED_BATCH_SIZE):
        batch = chunks[i : i + EMBED_BATCH_SIZE]
        all_vecs.append(embedder.embed(batch))
    vectors = np.vstack(all_vecs)

    doc = Document(dataset_id=dataset.id, filename=filename, size=path.stat().st_size, chunk_count=len(chunks))
    db.add(doc)
    db.flush()
    for i, (chunk_text, vec) in enumerate(zip(chunks, vectors)):
        db.add(
            Chunk(
                dataset_id=dataset.id,
                document_id=doc.id,
                idx=i,
                content=chunk_text,
                embedding=vec.astype(np.float32).tobytes(),
                dim=int(vec.shape[0]),
            )
        )
    db.commit()
    vector_store.invalidate(dataset.id)
    return doc


def ingest_table(
    db: Session, settings: Settings, dataset: Dataset, path: Path, filename: str
) -> dict:
    """CSV / Excel -> pandas -> SQLite 表，返回行列概要。"""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            df = pd.read_csv(path)
        except UnicodeDecodeError:
            df = pd.read_csv(path, encoding="gbk")
    elif suffix in (".xlsx", ".xls"):
        df = pd.read_excel(path)
    else:
        raise ValueError(f"表格数据集仅支持 csv/xlsx，收到: {suffix}")

    df = df.head(100000)  # 演示规模上限
    table_name = f"tbl_{uuid.uuid4().hex[:8]}"
    df.columns = [str(c).strip() for c in df.columns]
    # 先提交会话释放写锁，再用独立连接建表写入
    db.commit()
    df.to_sql(table_name, db.get_bind(), index=False, if_exists="fail")

    dataset.table_name = table_name
    dataset.meta = {
        "columns": list(df.columns),
        "dtypes": {c: str(t) for c, t in df.dtypes.items()},
        "rows": int(len(df)),
        "filename": filename,
    }
    doc = Document(dataset_id=dataset.id, filename=filename, size=path.stat().st_size, chunk_count=0, status="ready")
    db.add(doc)
    db.commit()
    return {"table_name": table_name, "rows": int(len(df)), "columns": list(df.columns)}
